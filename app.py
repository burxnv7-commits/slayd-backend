import os, io, json, re
import requests
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pptx import Presentation
from pptx.util import Inches, Pt

app=Flask(__name__); CORS(app)
API_KEY=os.environ.get('OPENROUTER_API_KEY','')
MODEL=os.environ.get('OPENROUTER_MODEL','openrouter/free')

def clean_json(s):
    s=re.sub(r'^```(?:json)?\s*|\s*```$','',s.strip(),flags=re.I)
    a=s.find('{'); b=s.rfind('}')
    return json.loads(s[a:b+1])

@app.get('/')
def home(): return jsonify(status='ok', service='Slayd Backend')

@app.post('/generate')
def generate():
    if not API_KEY: return jsonify(error='OPENROUTER_API_KEY serverda sozlanmagan'),500
    d=request.get_json(force=True) or {}
    topic=str(d.get('topic','')).strip(); n=max(3,min(int(d.get('slides',10)),25))
    lang=d.get('language','O‘zbekcha'); style=d.get('style','Professional')
    if not topic: return jsonify(error='Mavzu kiritilmagan'),400
    prompt=f'''{topic} mavzusida {n} slayddan iborat {lang} tilidagi {style} taqdimot tuz. Faqat JSON qaytar: {{"title":"...","slides":[{{"title":"...","bullets":["...","...","..."]}}]}}. Aynan {n} slayd bo'lsin. Har slaydda 3-5 qisqa, mazmunli punkt. Birinchi slayd titul, oxirgisi xulosa.'''
    r=requests.post('https://openrouter.ai/api/v1/chat/completions',headers={'Authorization':f'Bearer {API_KEY}','Content-Type':'application/json'},json={'model':MODEL,'messages':[{'role':'user','content':prompt}]},timeout=90)
    if not r.ok: return jsonify(error='AI xatosi',details=r.text[:800]),502
    try: data=clean_json(r.json()['choices'][0]['message']['content'])
    except Exception as e: return jsonify(error='AI javobini o‘qib bo‘lmadi',details=str(e)),502
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    for i,s in enumerate(data.get('slides',[])[:n]):
        layout=prs.slide_layouts[0] if i==0 else prs.slide_layouts[1]
        slide=prs.slides.add_slide(layout); slide.shapes.title.text=s.get('title','')
        if i==0:
            slide.placeholders[1].text=f'{topic}\nSlayd AI'
        else:
            tf=slide.placeholders[1].text_frame; tf.clear()
            for j,b in enumerate(s.get('bullets',[])[:5]):
                p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=str(b); p.font.size=Pt(24)
    out=io.BytesIO(); prs.save(out); out.seek(0)
    filename=re.sub(r'[^\w\-]+','_',topic,flags=re.UNICODE)[:50] or 'taqdimot'
    return send_file(out,as_attachment=True,download_name=f'{filename}.pptx',mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    return send_file(out,as_attachment=True,download_name=f'{filename}.pptx',mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')

@app.post("/webhooks/tezcheck")
def tezcheck_webhook():
    data = request.get_json(silent=True) or {}
    print("Tezcheck webhook:", data)
    return jsonify({"ok": True}), 200
